"""
Rapport de recherche : génération PDF / Word et dialog Streamlit.
"""
import pandas as pd
import streamlit as st

from utils import OP_NATURAL, _find_col, _date_label


def _report_data(df, conditions, current_table, joins, last_where: str) -> dict:
    from datetime import datetime as _dt
    cond_texts = []
    for c in conditions:
        op = OP_NATURAL.get(c.get("operator", ""), c.get("operator", ""))
        if c.get("is_bulk"):
            vals = ", ".join(c["values"][:5])
            sfx  = f" +{len(c['values'])-5} autres" if len(c["values"]) > 5 else ""
            cond_texts.append(f"{c['column']} {op} [{vals}{sfx}]")
        elif c.get("is_date"):
            cond_texts.append(f"{c['column']} en {_date_label(c['value'])}")
        else:
            cond_texts.append(f"{c['column']} {op} «{c['value']}»")

    join_texts = [f"{j['type']} {j['table']} ON {j['on']}" for j in (joins or [])]

    num_cols = df.select_dtypes(include=["number"]).columns.tolist()
    stats: dict = {}
    for col in num_cols:
        s = df[col].dropna()
        if len(s):
            stats[col] = {"min": float(s.min()), "max": float(s.max()),
                          "mean": float(s.mean()), "sum": float(s.sum()),
                          "count": int(s.count())}
    return {
        "title":        "Rapport d'analyse — SQL Query Builder",
        "generated_at": _dt.now().strftime("%d/%m/%Y à %H:%M"),
        "table":        current_table,
        "joins":        join_texts,
        "conditions":   cond_texts,
        "df":           df,
        "n_rows":       len(df),
        "n_cols":       len(df.columns),
        "stats":        stats,
    }


def generate_report_pdf(rd: dict, max_rows: int = 500) -> bytes:
    import io
    from reportlab.lib.pagesizes import A4
    from reportlab.lib             import colors
    from reportlab.lib.styles      import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units       import cm
    from reportlab.platypus        import (SimpleDocTemplate, Paragraph, Spacer,
                                           Table, TableStyle, HRFlowable)

    buf    = io.BytesIO()
    W, _H  = A4
    margin = 2.2 * cm
    doc    = SimpleDocTemplate(buf, pagesize=A4,
                               leftMargin=margin, rightMargin=margin,
                               topMargin=2.5*cm, bottomMargin=2*cm)
    cw     = W - 2 * margin

    styles  = getSampleStyleSheet()
    BLUE    = colors.HexColor("#1d4ed8")
    LBLUE   = colors.HexColor("#eff6ff")
    LGRAY   = colors.HexColor("#f8fafc")
    GRAY    = colors.HexColor("#e2e8f0")
    DARK    = colors.HexColor("#1e293b")
    MUTED   = colors.HexColor("#6b7280")

    S_TITLE  = ParagraphStyle("T",  parent=styles["Title"],    fontSize=17, spaceAfter=2,
                               textColor=colors.HexColor("#0f172a"))
    S_SUB    = ParagraphStyle("Su", parent=styles["Normal"],   fontSize=9,  spaceAfter=12,
                               textColor=MUTED)
    S_H1     = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=11, spaceBefore=14,
                               spaceAfter=5, textColor=colors.HexColor("#1e3a5f"))
    S_BODY   = ParagraphStyle("B",  parent=styles["Normal"],   fontSize=9,  leading=14,
                               spaceAfter=3, textColor=DARK)
    S_ITALIC = ParagraphStyle("I",  parent=S_BODY, textColor=MUTED)

    def _tbl(data, col_widths, header_bg=BLUE, alt=True):
        t = Table(data, colWidths=col_widths, repeatRows=1)
        cmds = [
            ("FONTNAME",      (0,0), (-1,0),  "Helvetica-Bold"),
            ("FONTNAME",      (0,1), (-1,-1), "Helvetica"),
            ("FONTSIZE",      (0,0), (-1,-1), 8),
            ("BACKGROUND",    (0,0), (-1,0),  header_bg),
            ("TEXTCOLOR",     (0,0), (-1,0),  colors.white),
            ("GRID",          (0,0), (-1,-1), 0.3, GRAY),
            ("TOPPADDING",    (0,0), (-1,-1), 4),
            ("BOTTOMPADDING", (0,0), (-1,-1), 4),
            ("LEFTPADDING",   (0,0), (-1,-1), 5),
            ("VALIGN",        (0,0), (-1,-1), "MIDDLE"),
        ]
        if alt:
            for i in range(1, len(data)):
                bg = LGRAY if i % 2 == 0 else colors.white
                cmds.append(("BACKGROUND", (0,i), (-1,i), bg))
        t.setStyle(TableStyle(cmds))
        return t

    story = []
    story.append(Paragraph(rd["title"], S_TITLE))
    story.append(Paragraph(f"Généré le {rd['generated_at']}", S_SUB))
    story.append(HRFlowable(width="100%", thickness=2, color=BLUE, spaceAfter=12))

    story.append(Paragraph("Paramètres de la requête", S_H1))
    story.append(Paragraph(f"<b>Table source :</b>  {rd['table']}", S_BODY))
    for j in rd["joins"]:
        story.append(Paragraph(f"<b>Jointure :</b>  {j}", S_BODY))
    if rd["conditions"]:
        story.append(Spacer(1, 4))
        story.append(Paragraph("<b>Conditions :</b>", S_BODY))
        for i, c in enumerate(rd["conditions"], 1):
            story.append(Paragraph(f" {i}.  {c}", S_BODY))
    else:
        story.append(Paragraph("<i>Aucune condition — tous les enregistrements.</i>", S_ITALIC))

    story.append(Paragraph("Résumé", S_H1))
    smry = [["Lignes retournées", str(rd["n_rows"])],
            ["Colonnes",          str(rd["n_cols"])],
            ["Table source",      rd["table"]],
            ["Généré le",         rd["generated_at"]]]
    story.append(_tbl([["Paramètre", "Valeur"]] + smry,
                      [cw * 0.4, cw * 0.6], header_bg=LBLUE))

    df2   = rd["df"].head(max_rows)
    shown = len(df2)
    sfx   = f" — {shown} affichées" if shown < rd["n_rows"] else ""
    story.append(Paragraph(f"Données  ({rd['n_rows']} lignes{sfx})", S_H1))
    cols   = list(df2.columns)
    n      = len(cols)
    cw_col = max(cw / n, 1.5*cm)
    rows   = [cols]
    for _, row in df2.iterrows():
        rows.append([("" if (v is None or str(v) == "nan") else str(v)) for v in row])
    story.append(_tbl(rows, [cw_col]*n))

    if rd["stats"]:
        story.append(Paragraph("Statistiques numériques", S_H1))
        hdr   = ["Colonne", "Min", "Max", "Moyenne", "Somme", "N"]
        srows = [hdr]
        for col, s in rd["stats"].items():
            srows.append([col, f"{s['min']:,.2f}", f"{s['max']:,.2f}",
                          f"{s['mean']:,.2f}", f"{s['sum']:,.2f}", str(s["count"])])
        story.append(_tbl(srows, [cw * 0.28] + [cw * 0.72 / 5] * 5,
                          header_bg=colors.HexColor("#0f172a")))

    df_full    = rd["df"]
    has_nom    = "nom"         in df_full.columns
    has_prenom = "prenom"      in df_full.columns
    has_dest   = "destination" in df_full.columns
    _id_col    = _find_col(df_full, "id", "clients_id")
    _stat_col  = _find_col(df_full, "statut", "clients_statut")

    if has_nom and has_prenom:
        from reportlab.platypus import KeepTogether
        S_NAME = ParagraphStyle("Name", parent=styles["Normal"],
                                 fontSize=10, fontName="Helvetica-Bold",
                                 textColor=colors.HexColor("#0f172a"),
                                 spaceBefore=10, spaceAfter=2)
        S_META = ParagraphStyle("Meta", parent=styles["Normal"],
                                 fontSize=8.5, leading=13,
                                 textColor=colors.HexColor("#6b7280"), spaceAfter=2)
        S_VG   = ParagraphStyle("Vg", parent=styles["Normal"],
                                 fontSize=8.5, leading=14,
                                 textColor=colors.HexColor("#1e293b"),
                                 leftIndent=14, spaceAfter=1)

        CONT_C = {"Asie": "#b45309", "Europe": "#1d4ed8", "Amérique": "#16a34a",
                  "Afrique": "#dc2626", "Océanie": "#7c3aed"}

        story.append(Spacer(1, 10))
        story.append(HRFlowable(width="100%", thickness=1,
                                 color=colors.HexColor("#e2e8f0"), spaceAfter=4))
        story.append(Paragraph("Fiches clients", S_H1))

        if has_dest and _id_col:
            for cid, grp in df_full.groupby(_id_col, sort=False):
                r0     = grp.iloc[0]
                prenom = r0.get("prenom", ""); nom = r0.get("nom", "")
                ville  = r0.get("ville", "");  email = r0.get("email", "")
                statut = r0.get(_stat_col) if _stat_col else r0.get("statut", "")
                s_str  = "actif" if statut == "actif" else "inactif"
                s_col  = "#16a34a" if statut == "actif" else "#dc2626"
                n_v    = len(grp)

                block = [Paragraph(
                    f"<b>{prenom} {nom}</b>"
                    f"<font size='8' color='#6b7280'>  ·  {ville}"
                    f"{'  ·  ' + email if email else ''}</font>"
                    f"  <font size='8' color='{s_col}'>{s_str}</font>"
                    f"  <font size='8' color='#6366f1'>{n_v} voyage(s)</font>",
                    S_NAME)]

                for _, vr in grp.iterrows():
                    dest  = vr.get("destination", ""); pays = vr.get("pays_destination", "")
                    cont  = vr.get("continent", "");  d_dep = str(vr.get("date_depart",""))[:10]
                    d_ret = str(vr.get("date_retour",""))[:10]
                    duree = vr.get("duree_jours"); tv = vr.get("type_voyage", "")
                    budget = vr.get("budget"); note = vr.get("note")
                    c_col  = CONT_C.get(cont, "#6b7280")
                    stars  = "★" * int(note) if note and not pd.isna(note) else "—"
                    bgt    = f"{int(budget):,}€" if budget and not pd.isna(budget) else "—"
                    duree_s = f"  ·  {int(duree)}j" if duree and not pd.isna(duree) else ""
                    block.append(Paragraph(
                        f"<font color='{c_col}'>▸</font>  "
                        f"<b>{dest}</b>"
                        f"<font color='#94a3b8'>  {pays}{duree_s}  ·  {d_dep} → {d_ret}"
                        f"  ·  {tv}</font>"
                        f"  <font color='#16a34a'>{bgt}</font>"
                        f"  <font color='#b45309' size='8'>{stars}</font>",
                        S_VG))

                block.append(HRFlowable(width="100%", thickness=0.5,
                                         color=colors.HexColor("#f1f5f9"), spaceAfter=2))
                story.append(KeepTogether(block))
        else:
            for _, r in df_full.iterrows():
                prenom = r.get("prenom", ""); nom = r.get("nom", "")
                ville  = r.get("ville", "");  email = r.get("email", "")
                tel    = r.get("telephone", ""); di = r.get("date_inscription", "")
                statut = r.get(_stat_col) if _stat_col else r.get("statut", "")
                s_col  = "#16a34a" if statut == "actif" else "#dc2626"
                story.append(Paragraph(
                    f"<b>{prenom} {nom}</b>"
                    f"<font size='8' color='#6b7280'>  ·  {ville}</font>"
                    f"  <font size='8' color='{s_col}'>{statut}</font>",
                    S_NAME))
                meta_parts = []
                if email: meta_parts.append(email)
                if tel:   meta_parts.append(tel)
                if di:    meta_parts.append(f"Membre depuis {str(di)[:10]}")
                if meta_parts:
                    story.append(Paragraph("  ·  ".join(meta_parts), S_META))

    doc.build(story)
    return buf.getvalue()


def generate_report_docx(rd: dict, max_rows: int = 500) -> bytes:
    import io
    from docx             import Document
    from docx.shared      import Pt, RGBColor, Cm
    from docx.enum.table  import WD_TABLE_ALIGNMENT
    from docx.oxml.ns     import qn
    from docx.oxml        import OxmlElement

    doc = Document()
    sec = doc.sections[0]
    from docx.shared import Cm as _Cm
    sec.page_width  = _Cm(21); sec.page_height = _Cm(29.7)
    sec.left_margin = sec.right_margin  = _Cm(2.5)
    sec.top_margin  = sec.bottom_margin = _Cm(2)
    content_cm = 16

    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)

    def _rgb(hex6):
        return RGBColor(int(hex6[:2],16), int(hex6[2:4],16), int(hex6[4:],16))

    def _shd(cell, fill_hex):
        tc = cell._tc; tcPr = tc.get_or_add_tcPr()
        for old in tcPr.findall(qn("w:shd")): tcPr.remove(old)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"),"clear"); shd.set(qn("w:color"),"auto")
        shd.set(qn("w:fill"), fill_hex.upper()); tcPr.append(shd)

    def _add_heading(text, level=1, color="0f172a"):
        p = doc.add_paragraph()
        r = p.add_run(text); r.bold = True
        r.font.size = Pt(14 if level == 1 else 11)
        r.font.color.rgb = _rgb(color)
        p.paragraph_format.space_before = Pt(14 if level == 1 else 8)
        p.paragraph_format.space_after  = Pt(5)

    def _add_kv(key, val):
        p = doc.add_paragraph()
        rk = p.add_run(f"{key} : "); rk.bold = True; rk.font.size = Pt(9)
        rv = p.add_run(str(val));    rv.font.size = Pt(9)
        p.paragraph_format.space_after = Pt(2)

    def _hr():
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(8)
        pPr = p._p.get_or_add_pPr(); pBdr = OxmlElement("w:pBdr")
        btm = OxmlElement("w:bottom")
        btm.set(qn("w:val"),"single"); btm.set(qn("w:sz"),"12")
        btm.set(qn("w:space"),"1");    btm.set(qn("w:color"),"1D4ED8")
        pBdr.append(btm); pPr.append(pBdr)

    def _tbl_docx(header_row, data_rows, col_widths_cm,
                  hdr_fill="1D4ED8", hdr_text="FFFFFF"):
        n = len(header_row); t = doc.add_table(rows=1, cols=n)
        t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.LEFT
        hcells = t.rows[0].cells
        for i, h in enumerate(header_row):
            hcells[i].text = str(h)
            hcells[i].width = Cm(col_widths_cm[i] if i < len(col_widths_cm) else 2)
            r2 = hcells[i].paragraphs[0].runs[0]
            r2.bold = True; r2.font.size = Pt(8)
            r2.font.color.rgb = _rgb(hdr_text); _shd(hcells[i], hdr_fill)
        for ri, row in enumerate(data_rows):
            cells = t.add_row().cells; fill = "F8FAFC" if ri % 2 == 0 else "FFFFFF"
            for i, v in enumerate(row):
                cells[i].text  = str(v)
                cells[i].width = Cm(col_widths_cm[i] if i < len(col_widths_cm) else 2)
                for rx in cells[i].paragraphs[0].runs: rx.font.size = Pt(8)
                _shd(cells[i], fill)
        return t

    p = doc.add_paragraph()
    r = p.add_run(rd["title"]); r.bold = True; r.font.size = Pt(17)
    r.font.color.rgb = _rgb("0f172a"); p.paragraph_format.space_after = Pt(2)
    p2 = doc.add_paragraph(f"Généré le {rd['generated_at']}")
    p2.runs[0].font.size = Pt(9); p2.runs[0].font.color.rgb = _rgb("6b7280")
    p2.paragraph_format.space_after = Pt(4); _hr()

    _add_heading("Paramètres de la requête")
    _add_kv("Table source", rd["table"])
    for j in rd["joins"]: _add_kv("Jointure", j)
    if rd["conditions"]:
        p_h = doc.add_paragraph()
        rh = p_h.add_run("Conditions :"); rh.bold = True; rh.font.size = Pt(9)
        p_h.paragraph_format.space_after = Pt(2)
        for i, c in enumerate(rd["conditions"], 1):
            pb = doc.add_paragraph(f"  {i}.  {c}", style="List Bullet")
            pb.paragraph_format.space_after = Pt(1)
            for rx in pb.runs: rx.font.size = Pt(9)
    else:
        p_nc = doc.add_paragraph("Aucune condition — tous les enregistrements.")
        p_nc.runs[0].italic = True; p_nc.runs[0].font.size = Pt(9)

    _add_heading("Résumé")
    smry = [["Lignes retournées", str(rd["n_rows"])],
            ["Colonnes",          str(rd["n_cols"])],
            ["Table source",      rd["table"]],
            ["Généré le",         rd["generated_at"]]]
    _tbl_docx(["Paramètre", "Valeur"], smry,
              [content_cm * 0.4, content_cm * 0.6],
              hdr_fill="DBEAFE", hdr_text="1E3A5F")

    df2   = rd["df"].head(max_rows)
    shown = len(df2); sfx = f" — {shown} affichées" if shown < rd["n_rows"] else ""
    _add_heading(f"Données  ({rd['n_rows']} lignes{sfx})")
    cols = list(df2.columns); n = len(cols); cw = max(content_cm / n, 1.5)
    rows = [["" if (v is None or str(v) == "nan") else str(v) for v in row]
            for _, row in df2.iterrows()]
    _tbl_docx(cols, rows, [cw]*n)

    if rd["stats"]:
        _add_heading("Statistiques numériques")
        hdr   = ["Colonne", "Min", "Max", "Moyenne", "Somme", "N"]
        sdata = [[c, f"{s['min']:,.2f}", f"{s['max']:,.2f}",
                  f"{s['mean']:,.2f}", f"{s['sum']:,.2f}", str(s["count"])]
                 for c, s in rd["stats"].items()]
        _tbl_docx(hdr, sdata, [content_cm * 0.28] + [content_cm * 0.72 / 5] * 5,
                  hdr_fill="0F172A")

    df_full    = rd["df"]
    has_nom    = "nom"         in df_full.columns
    has_prenom = "prenom"      in df_full.columns
    has_dest   = "destination" in df_full.columns
    _id_col    = _find_col(df_full, "id", "clients_id")
    _stat_col  = _find_col(df_full, "statut", "clients_statut")

    if has_nom and has_prenom:
        from docx.oxml import OxmlElement as _OxmlElement
        p_hr = doc.add_paragraph()
        p_hr.paragraph_format.space_before = Pt(10)
        p_hr.paragraph_format.space_after  = Pt(4)
        pPr2  = p_hr._p.get_or_add_pPr(); pBdr2 = _OxmlElement("w:pBdr")
        btm2  = _OxmlElement("w:bottom")
        btm2.set(qn("w:val"),"single"); btm2.set(qn("w:sz"),"4")
        btm2.set(qn("w:space"),"1");    btm2.set(qn("w:color"),"E2E8F0")
        pBdr2.append(btm2); pPr2.append(pBdr2)
        _add_heading("Fiches clients")

        CONT_C = {"Asie": "B45309", "Europe": "1D4ED8", "Amérique": "16A34A",
                  "Afrique": "DC2626", "Océanie": "7C3AED"}

        def _run(para, text, size=9, bold=False, color="1e293b", italic=False):
            r = para.add_run(text); r.bold = bold; r.italic = italic
            r.font.size = Pt(size); r.font.color.rgb = _rgb(color.lstrip("#"))
            return r

        if has_dest and _id_col:
            for cid, grp in df_full.groupby(_id_col, sort=False):
                r0     = grp.iloc[0]
                prenom = r0.get("prenom",""); nom = r0.get("nom","")
                ville  = r0.get("ville","");  email = r0.get("email","")
                statut = r0.get(_stat_col) if _stat_col else r0.get("statut","")
                s_col  = "16A34A" if statut == "actif" else "DC2626"
                n_v    = len(grp)
                ph = doc.add_paragraph()
                ph.paragraph_format.space_before = Pt(10)
                ph.paragraph_format.space_after  = Pt(2)
                _run(ph, f"{prenom} {nom}", size=10, bold=True, color="0f172a")
                _run(ph, f"   {ville}", size=8.5, color="6b7280")
                if email: _run(ph, f"  ·  {email}", size=8.5, color="6b7280")
                _run(ph, f"   {'actif' if statut == 'actif' else 'inactif'}",
                     size=8.5, color=s_col)
                _run(ph, f"  ·  {n_v} voyage(s)", size=8.5, color="6366f1")
                for _, vr in grp.iterrows():
                    dest  = vr.get("destination",""); pays = vr.get("pays_destination","")
                    cont  = vr.get("continent","");  d_dep = str(vr.get("date_depart",""))[:10]
                    d_ret = str(vr.get("date_retour",""))[:10]
                    duree = vr.get("duree_jours"); tv = vr.get("type_voyage","")
                    budget = vr.get("budget"); note = vr.get("note")
                    c_col  = CONT_C.get(cont, "6b7280")
                    stars  = "★" * int(note) if note and not pd.isna(note) else "—"
                    bgt    = f"{int(budget):,}€" if budget and not pd.isna(budget) else "—"
                    duree_s = f"  ·  {int(duree)}j" if duree and not pd.isna(duree) else ""
                    pv = doc.add_paragraph()
                    pv.paragraph_format.space_after = Pt(1)
                    pv.paragraph_format.left_indent = Pt(14)
                    _run(pv, "▸  ", size=9, color=c_col, bold=True)
                    _run(pv, dest, size=9, bold=True, color="1e3a5f")
                    _run(pv, f"   {pays}{duree_s}  ·  {d_dep} → {d_ret}  ·  {tv}",
                         size=8, color="94a3b8")
                    _run(pv, f"   {bgt}", size=9, color="16a34a")
                    _run(pv, f"  {stars}", size=8, color="b45309")
                p_sep = doc.add_paragraph()
                p_sep.paragraph_format.space_before = Pt(4)
                p_sep.paragraph_format.space_after  = Pt(0)
                pPr3  = p_sep._p.get_or_add_pPr(); pBdr3 = _OxmlElement("w:pBdr")
                b3 = _OxmlElement("w:bottom")
                b3.set(qn("w:val"),"single"); b3.set(qn("w:sz"),"2")
                b3.set(qn("w:space"),"1");    b3.set(qn("w:color"),"F1F5F9")
                pBdr3.append(b3); pPr3.append(pBdr3)
        else:
            for _, r in df_full.iterrows():
                prenom = r.get("prenom",""); nom = r.get("nom","")
                ville  = r.get("ville","");  email = r.get("email","")
                tel    = r.get("telephone",""); di = r.get("date_inscription","")
                statut = r.get(_stat_col) if _stat_col else r.get("statut","")
                s_col  = "16A34A" if statut == "actif" else "DC2626"
                pc = doc.add_paragraph()
                pc.paragraph_format.space_before = Pt(8)
                pc.paragraph_format.space_after  = Pt(2)
                _run(pc, f"{prenom} {nom}", size=10, bold=True, color="0f172a")
                _run(pc, f"  ·  {ville}", size=8.5, color="6b7280")
                _run(pc, f"  ·  {'actif' if statut == 'actif' else 'inactif'}",
                     size=8.5, color=s_col)
                meta = []
                if email: meta.append(email)
                if tel:   meta.append(tel)
                if di:    meta.append(f"Membre depuis {str(di)[:10]}")
                if meta:
                    pm = doc.add_paragraph("   ·   ".join(meta))
                    pm.runs[0].font.size = Pt(8.5)
                    pm.runs[0].font.color.rgb = _rgb("6b7280")
                    pm.paragraph_format.space_after = Pt(2)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@st.dialog("📄 Générer un rapport")
def _rapport_dialog(df, conditions, current_table, joins, last_where):
    st.markdown(
        f"<p style='color:#94a3b8;font-size:.85rem;'>"
        f"{len(df)} lignes · {len(df.columns)} colonnes · table <b>{current_table}</b></p>",
        unsafe_allow_html=True)

    c1, c2 = st.columns(2)
    fmt    = c1.selectbox("Format", ["PDF", "Word (.docx)"], key="rpt_fmt")
    max_r  = c2.number_input(
        "Lignes max.", min_value=1, max_value=len(df),
        value=min(len(df), 500), step=max(1, min(50, len(df))), key="rpt_maxrows",
    )

    if st.button("⬇ Générer", type="primary", width="stretch", key="rpt_gen"):
        with st.spinner("Génération en cours…"):
            rd = _report_data(df, conditions, current_table, joins, last_where)
            try:
                if fmt == "PDF":
                    data  = generate_report_pdf(rd, max_rows=int(max_r))
                    fname = f"rapport_{current_table}.pdf"
                    mime  = "application/pdf"
                else:
                    data  = generate_report_docx(rd, max_rows=int(max_r))
                    fname = f"rapport_{current_table}.docx"
                    mime  = ("application/vnd.openxmlformats-officedocument"
                             ".wordprocessingml.document")
                st.download_button(
                    f"⬇  Télécharger  {fname}",
                    data=data, file_name=fname, mime=mime,
                    key="rpt_dl", use_container_width=True, type="primary",
                )
            except ImportError as e:
                st.error(
                    f"Bibliothèque manquante : `{e}`\n\n"
                    f"Ajoutez `{'reportlab' if fmt == 'PDF' else 'python-docx'}` "
                    f"à votre `requirements.txt` et redémarrez l'app.")
            except Exception as e:
                st.error(f"Erreur de génération : {e}")
