"""
Rapport des mutations du personnel : génération PDF / Word et dialog Streamlit.
"""
import pandas as pd
import streamlit as st


def generate_report_mutations_pdf(data: dict, year_from: int, year_to: int) -> bytes:
    import io
    from reportlab.lib.pagesizes import A4
    from reportlab.lib             import colors
    from reportlab.lib.styles      import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units       import cm
    from reportlab.platypus        import (SimpleDocTemplate, Paragraph, Spacer,
                                            Table, TableStyle, HRFlowable, PageBreak)
    from datetime import datetime as _dt

    mut   = data["mutations"]
    cur   = data["current"]
    buf   = io.BytesIO()
    W, _H = A4
    M     = 2.2 * cm
    doc   = SimpleDocTemplate(buf, pagesize=A4,
                               leftMargin=M, rightMargin=M,
                               topMargin=2.5*cm, bottomMargin=2*cm)
    cw    = W - 2 * M

    styles = getSampleStyleSheet()
    BLUE   = colors.HexColor("#1d4ed8")
    LGRAY  = colors.HexColor("#f8fafc")
    GRAY   = colors.HexColor("#e2e8f0")
    DARK   = colors.HexColor("#0f172a")
    MUTED  = colors.HexColor("#6b7280")

    S_TITLE = ParagraphStyle("T",  parent=styles["Title"],   fontSize=17, spaceAfter=2,
                              textColor=DARK)
    S_SUB   = ParagraphStyle("Su", parent=styles["Normal"],  fontSize=9,  spaceAfter=12,
                              textColor=MUTED)
    S_H1    = ParagraphStyle("H1", parent=styles["Heading1"],fontSize=12, spaceBefore=14,
                              spaceAfter=6, textColor=colors.HexColor("#1e3a5f"))
    S_BODY  = ParagraphStyle("B",  parent=styles["Normal"],  fontSize=9,  leading=14,
                              spaceAfter=3)

    def _tbl(data, cws, hbg=BLUE):
        t = Table(data, colWidths=cws, repeatRows=1)
        cmds = [
            ("FONTNAME",      (0,0), (-1,0),  "Helvetica-Bold"),
            ("FONTNAME",      (0,1), (-1,-1), "Helvetica"),
            ("FONTSIZE",      (0,0), (-1,-1), 8),
            ("BACKGROUND",    (0,0), (-1,0),  hbg),
            ("TEXTCOLOR",     (0,0), (-1,0),  colors.white),
            ("GRID",          (0,0), (-1,-1), 0.3, GRAY),
            ("TOPPADDING",    (0,0), (-1,-1), 4),
            ("BOTTOMPADDING", (0,0), (-1,-1), 4),
            ("LEFTPADDING",   (0,0), (-1,-1), 5),
        ]
        for i in range(1, len(data)):
            cmds.append(("BACKGROUND", (0,i), (-1,i), LGRAY if i%2==0 else colors.white))
        t.setStyle(TableStyle(cmds))
        return t

    if not mut.empty and "annee" in mut.columns:
        mut_p = mut[(mut["annee"] >= year_from) & (mut["annee"] <= year_to)].copy()
    else:
        mut_p = pd.DataFrame()

    story = []
    story.append(Paragraph("Rapport des Mutations du Personnel", S_TITLE))
    story.append(Paragraph(
        f"Agence de voyages  ·  Période {year_from}–{year_to}  ·  "
        f"Généré le {_dt.now().strftime('%d/%m/%Y')}", S_SUB))
    story.append(HRFlowable(width="100%", thickness=2, color=BLUE, spaceAfter=14))

    story.append(Paragraph("Résumé exécutif", S_H1))
    n_mut      = len(mut_p)
    n_employes = mut_p["employe_id"].nunique() if not mut_p.empty else 0
    n_pays_imp = len(set(list(mut_p.get("pays_depart", pd.Series()).unique()) +
                         list(mut_p.get("pays_arrivee", pd.Series()).unique()))) if not mut_p.empty else 0
    n_actifs   = len(cur)
    smry = [["Indicateur", "Valeur"],
            ["Employés actifs (postes actuels)",   str(n_actifs)],
            ["Mutations sur la période",           str(n_mut)],
            ["Employés ayant muté",                str(n_employes)],
            ["Pays impliqués",                     str(n_pays_imp)],
            ["Moyenne mutations / an",
             f"{n_mut / max(year_to - year_from + 1, 1):.1f}"]]
    story.append(_tbl(smry, [cw*0.55, cw*0.45], hbg=colors.HexColor("#dbeafe")))
    story.append(Spacer(1, 8))
    story.append(Paragraph(
        f"Sur la période {year_from}–{year_to}, l'agence a enregistré "
        f"<b>{n_mut} mutation(s)</b> impliquant <b>{n_employes} employé(s)</b> "
        f"dans <b>{n_pays_imp} pays</b>.", S_BODY))

    story.append(Paragraph("Analyse par année", S_H1))
    if not mut_p.empty:
        by_year = (mut_p.groupby("annee")
                        .agg(nb_mutations=("employe_id","count"),
                             nb_employes=("employe_id","nunique"),
                             pays_dest=("pays_arrivee","nunique"))
                        .reset_index().sort_values("annee"))
        hdr  = ["Année", "Mutations", "Employés mutés", "Pays de destination"]
        rows = [hdr] + [[str(int(r["annee"])), str(r["nb_mutations"]),
                          str(r["nb_employes"]), str(r["pays_dest"])]
                         for _, r in by_year.iterrows()]
        story.append(_tbl(rows, [cw*0.2, cw*0.25, cw*0.3, cw*0.25]))
    else:
        story.append(Paragraph("Aucune mutation sur la période sélectionnée.", S_BODY))

    story.append(Paragraph("Pays d'accueil les plus fréquents", S_H1))
    if not mut_p.empty and "pays_arrivee" in mut_p.columns:
        top_pays = mut_p["pays_arrivee"].value_counts().head(8)
        rows2 = [["Pays d'accueil", "Nb mutations", "Employés"]]
        for pays, nb in top_pays.items():
            emp = mut_p[mut_p["pays_arrivee"]==pays]["employe_id"].nunique()
            rows2.append([pays, str(nb), str(emp)])
        story.append(_tbl(rows2, [cw*0.5, cw*0.25, cw*0.25]))

    story.append(Paragraph("Employés les plus mobiles", S_H1))
    if not mut_p.empty:
        mob = (mut_p.groupby(["employe_id","nom","prenom","poste"])
                    .size().reset_index(name="nb_mutations")
                    .sort_values("nb_mutations", ascending=False).head(10))
        rows3 = [["Nom","Prénom","Poste","Mutations"]] + \
                [[r["nom"],r["prenom"],r["poste"],str(r["nb_mutations"])]
                 for _, r in mob.iterrows()]
        story.append(_tbl(rows3, [cw*0.22, cw*0.22, cw*0.38, cw*0.18]))

    story.append(PageBreak())
    story.append(Paragraph("Détail chronologique des mutations", S_H1))
    if not mut_p.empty:
        rows4 = [["Date","Employé","Poste","Départ","Arrivée"]]
        for _, r in mut_p.sort_values("date_mutation", ascending=False).iterrows():
            rows4.append([
                str(r.get("date_mutation",""))[:10],
                f"{r.get('prenom','')} {r.get('nom','')}",
                r.get("poste",""),
                f"{r.get('agence_depart','')} ({r.get('pays_depart','')})",
                f"{r.get('agence_arrivee','')} ({r.get('pays_arrivee','')})",
            ])
        story.append(_tbl(rows4, [cw*0.1, cw*0.2, cw*0.2, cw*0.25, cw*0.25]))

    doc.build(story)
    return buf.getvalue()


def generate_report_mutations_docx(data: dict, year_from: int, year_to: int) -> bytes:
    import io
    from docx            import Document
    from docx.shared     import Pt, Cm as _Cm, RGBColor
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns    import qn
    from docx.oxml       import OxmlElement
    from datetime import datetime as _dt

    mut   = data["mutations"]
    cur   = data["current"]
    doc   = Document()
    sec   = doc.sections[0]
    sec.page_width  = _Cm(21); sec.page_height = _Cm(29.7)
    sec.left_margin = sec.right_margin = _Cm(2.5)
    sec.top_margin  = sec.bottom_margin = _Cm(2)
    content_cm = 16

    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)

    def _rgb(h): return RGBColor(int(h[:2],16), int(h[2:4],16), int(h[4:],16))

    def _shd(cell, fill):
        tc = cell._tc; tcPr = tc.get_or_add_tcPr()
        for old in tcPr.findall(qn("w:shd")): tcPr.remove(old)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"),"clear"); shd.set(qn("w:color"),"auto")
        shd.set(qn("w:fill"), fill.upper()); tcPr.append(shd)

    def _heading(txt, color="0f172a", size=14):
        p = doc.add_paragraph(); r = p.add_run(txt)
        r.bold = True; r.font.size = Pt(size); r.font.color.rgb = _rgb(color)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after  = Pt(5)

    def _tbl(hdr, rows, cws, hfill="1D4ED8"):
        n = len(hdr); t = doc.add_table(rows=1, cols=n)
        t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.LEFT
        hcells = t.rows[0].cells
        for i, h in enumerate(hdr):
            hcells[i].text = str(h); hcells[i].width = _Cm(cws[i])
            r2 = hcells[i].paragraphs[0].runs[0]
            r2.bold = True; r2.font.size = Pt(8)
            r2.font.color.rgb = _rgb("FFFFFF"); _shd(hcells[i], hfill)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells; fill = "F8FAFC" if ri%2==0 else "FFFFFF"
            for i, v in enumerate(row):
                cells[i].text = str(v); cells[i].width = _Cm(cws[i])
                for rx in cells[i].paragraphs[0].runs: rx.font.size = Pt(8)
                _shd(cells[i], fill)
        return t

    if not mut.empty and "annee" in mut.columns:
        mut_p = mut[(mut["annee"] >= year_from) & (mut["annee"] <= year_to)].copy()
    else:
        mut_p = pd.DataFrame()

    p = doc.add_paragraph(); r = p.add_run("Rapport des Mutations du Personnel")
    r.bold = True; r.font.size = Pt(17); r.font.color.rgb = _rgb("0f172a")
    p.paragraph_format.space_after = Pt(2)
    p2 = doc.add_paragraph(
        f"Période {year_from}–{year_to}  ·  Généré le {_dt.now().strftime('%d/%m/%Y')}")
    p2.runs[0].font.size = Pt(9); p2.runs[0].font.color.rgb = _rgb("6b7280")
    p2.paragraph_format.space_after = Pt(6)

    _heading("Résumé exécutif")
    n_mut = len(mut_p); n_emp = mut_p["employe_id"].nunique() if not mut_p.empty else 0
    smry  = [["Employés actifs", str(len(cur))], ["Mutations", str(n_mut)],
             ["Employés mutés",  str(n_emp)],
             ["Moy. mutations/an", f"{n_mut/max(year_to-year_from+1,1):.1f}"]]
    _tbl(["Indicateur","Valeur"], smry,
         [content_cm*0.55, content_cm*0.45], hfill="DBEAFE")

    _heading("Par année")
    if not mut_p.empty:
        by_y = (mut_p.groupby("annee").agg(nb=("employe_id","count"),
                emp=("employe_id","nunique")).reset_index().sort_values("annee"))
        _tbl(["Année","Mutations","Employés mutés"],
             [[str(int(r["annee"])), str(r["nb"]), str(r["emp"])]
              for _, r in by_y.iterrows()],
             [content_cm*0.25, content_cm*0.4, content_cm*0.35])

    _heading("Pays d'accueil les plus fréquents")
    if not mut_p.empty and "pays_arrivee" in mut_p.columns:
        top = mut_p["pays_arrivee"].value_counts().head(8)
        _tbl(["Pays","Mutations"], [[p, str(n)] for p, n in top.items()],
             [content_cm*0.6, content_cm*0.4])

    _heading("Employés les plus mobiles")
    if not mut_p.empty:
        mob = (mut_p.groupby(["nom","prenom","poste"]).size()
                    .reset_index(name="n").sort_values("n", ascending=False).head(10))
        _tbl(["Nom","Prénom","Poste","Nb mutations"],
             [[r["nom"],r["prenom"],r["poste"],str(r["n"])] for _, r in mob.iterrows()],
             [content_cm*0.22, content_cm*0.22, content_cm*0.38, content_cm*0.18])

    _heading("Détail chronologique")
    if not mut_p.empty:
        _tbl(["Date","Employé","Départ","Arrivée"],
             [[str(r.get("date_mutation",""))[:10],
               f"{r.get('prenom','')} {r.get('nom','')}",
               f"{r.get('pays_depart','')}",
               f"{r.get('ville_arrivee','')} ({r.get('pays_arrivee','')})"]
              for _, r in mut_p.sort_values("date_mutation", ascending=False).iterrows()],
             [content_cm*0.12, content_cm*0.22, content_cm*0.33, content_cm*0.33])

    buf = io.BytesIO(); doc.save(buf)
    return buf.getvalue()


@st.dialog("📋 Rapport mutations du personnel")
def _mutations_rapport_dialog(data: dict) -> None:
    mut   = data.get("mutations", pd.DataFrame())
    years = (sorted(mut["annee"].dropna().unique().astype(int))
             if not mut.empty and "annee" in mut.columns else [2019, 2024])
    y_min, y_max = (int(min(years)), int(max(years))) if years else (2019, 2024)

    st.markdown(
        f"<p style='color:#94a3b8;font-size:.85rem;'>"
        f"{len(mut)} mutation(s)  ·  {y_min}–{y_max}</p>",
        unsafe_allow_html=True)

    c1, c2 = st.columns(2)
    fmt    = c1.selectbox("Format", ["PDF", "Word (.docx)"], key="mrpt_fmt")
    period = c2.select_slider(
        "Période", options=list(range(y_min, y_max+1)),
        value=(y_min, y_max), key="mrpt_period",
    )

    if st.button("⬇ Générer", type="primary", use_container_width=True, key="mrpt_gen"):
        with st.spinner("Génération…"):
            try:
                yf, yt = int(period[0]), int(period[1])
                if fmt == "PDF":
                    raw   = generate_report_mutations_pdf(data, yf, yt)
                    fname = f"mutations_personnel_{yf}_{yt}.pdf"
                    mime  = "application/pdf"
                else:
                    raw   = generate_report_mutations_docx(data, yf, yt)
                    fname = f"mutations_personnel_{yf}_{yt}.docx"
                    mime  = ("application/vnd.openxmlformats-officedocument"
                             ".wordprocessingml.document")
                st.download_button(f"⬇  {fname}", data=raw,
                                   file_name=fname, mime=mime,
                                   use_container_width=True, key="mrpt_dl")
            except Exception as e:
                st.error(f"Erreur : {e}")
